"""Build the revised report DOCX; equations use pre-rendered MathJax images."""
from pathlib import Path
import re,shutil,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'report'
source=(REPORT/'Folsom水库调度研究报告_完整版.md').read_text(encoding='utf-8')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.72)
sec.header_distance=sec.footer_distance=Inches(.28)
width=7.06

def font(style,size,east='宋体',western='Times New Roman'):
 s=doc.styles[style];s.font.name=western;s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)
 return s
normal=font('Normal',11);normal.paragraph_format.line_spacing=1.25
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
font('Title',18,'黑体').paragraph_format.space_after=Pt(10)
font('Subtitle',12,'宋体').paragraph_format.space_after=Pt(10)
for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
 s=font(name,size,'黑体');s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
font('Caption',9.5).paragraph_format.space_after=Pt(6)
font('Header',8);font('Footer',8)
sec.different_first_page_header_footer=True
sec.header.paragraphs[0].text='Folsom水库供水保障与防洪约束下的调度策略比较'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')

def link(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'175287');pr.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
 t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text):
 pattern=r'\[([^\]]+)\]\((https?://[^)]+)\)|(https?://[^\s。；，）]+)'
 end=0
 for m in re.finditer(pattern,text):
  p.add_run(text[end:m.start()].replace('**','').replace('`',''))
  if m.group(1):link(p,m.group(1),m.group(2))
  else:link(p,m.group(3),m.group(3))
  end=m.end()
 p.add_run(text[end:].replace('**','').replace('`',''))


# Independent cover and deterministic linked contents page.
for _ in range(4):doc.add_paragraph()
p=doc.add_paragraph('研究报告',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Folsom水库供水保障与防洪约束下的调度策略比较',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('基于官方水文输入与情景需水的可复现研究',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(4):doc.add_paragraph()
for text in ['研究对象  美国加州 Folsom 水库','研究范围  数据质量审查  调度策略比较  稳健性与数值验证','报告日期  2026年10月23日','版本  完整研究报告']:
 p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
p=doc.add_paragraph('目录',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
headings=re.findall(r'^## (.+)$',source,re.M)
pages_file=REPORT/'full_report_contents.json'
pages=json.loads(pages_file.read_text(encoding='utf-8')) if pages_file.exists() else {}
for j,h in enumerate(headings):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
 p.paragraph_format.tab_stops.add_tab_stop(Inches(6.8),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 hl=OxmlElement('w:hyperlink');hl.set(qn('w:anchor'),'chapter_'+str(j));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=h;r.append(t);hl.append(r);p._p.append(hl)
 p.add_run('\t'+str(pages.get(h,'待校准')))
doc.add_paragraph('目录页码对应本版完整文档页码。',style='Caption')
doc.add_page_break()
chapter_index=0

eqnum=0
lines=source.splitlines();i=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('|'):
  data=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[-: ]+',x) for x in cells):data.append(cells)
   i+=1
  n=len(data[0]);t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  if n==3:weights=[1.2,1.5,3.6]
  elif n==4:weights=[1.15,1.6,1.8,1.2]
  elif n==5:weights=[1.9,1.25,1.25,1.7,1.5]
  else:weights=[1.3]+[1.]*(n-1)
  widths=[width*w/sum(weights) for w in weights]
  for col,w in zip(t.columns,widths):col.width=Inches(w)
  props=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
  for side in ['top','left','bottom','right','insideH','insideV']:
   edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
  props.append(borders)
  for ri,row in enumerate(data):
   cells=t.add_row().cells
   if ri==0:
    repeat=OxmlElement('w:tblHeader');t.rows[-1]._tr.get_or_add_trPr().append(repeat)
   prevent=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(prevent)
   for j,(cell,value) in enumerate(zip(cells,row)):
    cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
    for side in ['top','bottom','left','right']:
     v=OxmlElement('w:'+side);v.set(qn('w:w'),'80');v.set(qn('w:type'),'dxa');marg.append(v)
    tcpr.append(marg)
    if ri==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE7EF');tcpr.append(shade)
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (ri==0 or (j>0 and n>=4)) else WD_ALIGN_PARAGRAPH.LEFT
    inline(p,value)
    for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
  doc.add_paragraph().paragraph_format.space_after=Pt(2)
  continue
 if line.startswith('!['):
  alt,path=re.match(r'!\[(.*?)\]\((.*?)\)',line).groups();p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.keep_with_next=True
  im=Image.open(REPORT/path);iw,ih=im.size
  max_height=3.8 if '05_parameter_search' in path else 5.0
  pw=min(width,max_height*iw/ih)
  run=p.add_run();shape=run.add_picture(str(REPORT/path),width=Inches(pw));shape._inline.docPr.set('descr',alt)
  cp=doc.add_paragraph(alt,style='Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif line.startswith('# '):pass
 elif line.startswith('## 基于'):doc.add_paragraph(line[3:],style='Subtitle')
 elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
 elif line.startswith('## '):
  p=doc.add_paragraph(line[3:],style='Heading 1')
  b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(chapter_index));b.set(qn('w:name'),'chapter_'+str(chapter_index));p._p.insert(0,b)
  e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(chapter_index));p._p.append(e);chapter_index+=1
 elif line.startswith('`') and line.endswith('`'):
  eqnum+=1;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  fp=REPORT/'equations'/f'eq{eqnum}.png';im=Image.open(fp)
  target=min(width-.2,im.width/300)
  p.add_run().add_picture(str(fp),width=Inches(target))
  p.paragraph_format.space_after=Pt(8)
 else:
  p=doc.add_paragraph();inline(p,line)
 i+=1
assert eqnum==7
# Remove inherited template borders, including the default blue title rule.
for element in [doc.styles.element,doc.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):
  for run in p.runs:run.text=run.text.replace('、',' ').replace('，',' ')
doc.core_properties.title='Folsom水库供水保障与防洪约束下的调度策略比较（数值修正版）'
doc.core_properties.subject='官方水文输入、数值时间步长修正及调度策略情景研究'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
path=REPORT/'Folsom水库调度研究报告_完整版.docx';doc.save(path)
print('Saved',path)
